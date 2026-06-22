import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

class ExcelGenerator:
    def export_data(self, filename, headers, data, title="Report"):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.sheet_view.rightToLeft = True

        # Style for header
        header_fill = PatternFill(start_color="1a2a6c", end_color="1a2a6c", fill_type="solid")
        header_font = Font(color="ffffff", bold=True)

        # Add Title
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        ws.cell(row=1, column=1, value=title).font = Font(size=14, bold=True)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")

        # Add Headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Add Data
        for row_idx, row_data in enumerate(data, 3):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value).alignment = Alignment(horizontal="center")

        # adjust column width
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except: pass
            ws.column_dimensions[column_letter].width = max_length + 2

        wb.save(filename)
        return filename

    def import_items(self, filename):
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        items = []
        # Expecting Code, Name, Category, Unit, MinStock
        for row in ws.iter_rows(min_row=3, values_only=True):
            if row[0]:
                items.append({
                    "code": row[0],
                    "name": row[1],
                    "category": row[2],
                    "unit": row[3],
                    "min_stock": row[4] or 0
                })
        return items
